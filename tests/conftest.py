"""pytest 共有フィクスチャ（当直くん回帰テスト）。

方針:
- main.run() は重い（num_patterns=20 で ~2分）。局所探索が支配的で num_patterns を
  絞っても大きくは縮まないため、**セッション全体で1回だけ** 実行し、生成された
  出力xlsxを全テストで共有する。
- 出力xlsxは openpyxl で「独立に」読み直し、制約充足を検証する。
- 期待値（TARGET_CAP / EXTRA_ALLOWED / 可否コード / 祝日 等）は run() 実行後に
  main モジュールへ公開されるグローバル（get_avail_code, TARGET_CAP, ...）を
  参照の真実源として用いる。出力の「値」はxlsx由来、「制約」はmain由来という分離。

決定性メモ:
  同一入力・num_patterns=20 の2回実行で pattern_01 は完全一致した（PYTHONHASHSEED
  未設定でも一致）。ただし将来のset反復順依存の非決定性リスクを避けるため、
  スナップショット比較ではなく不変条件テストを採用する。
"""

import os
import sys

import openpyxl
import pandas as pd
import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

INPUT_XLSX = os.path.join(REPO_ROOT, "docs", "Tochoku.v7_2026.03.1.xlsx")

# 出力が重いので num_patterns は環境変数で上書き可能（既定20）。
NUM_PATTERNS = int(os.environ.get("TOCHOKU_TEST_NUM_PATTERNS", "20"))


def _read_input_slots(main_mod):
    """入力 sheet1 から枠（スロット）情報を **main のパイプラインとは独立に** 再構築する。

    Returns:
        docs: 正規化済み医師名の set（Sheet2 のヘッダ由来）
        slots: {(row_pos, colname)}  枠として扱うセル（自動枠 or 事前割当）
        preassigned: {(row_pos, colname): 正規化医師名}
        dates: {row_pos: pd.Timestamp}  各行の日付（正規化）
        hospital_cols: 病院列名リスト（カテ当番を除外）
    """
    xls = pd.ExcelFile(INPUT_XLSX)
    sheet1 = main_mod.find_sheet_name(xls, "sheet1")
    sheet2 = main_mod.find_sheet_name(xls, "sheet2") or main_mod.find_sheet_name(xls, "Sheet2")

    s1 = main_mod.strip_cols(pd.read_excel(xls, sheet_name=sheet1))
    s1.columns = main_mod.make_unique(list(s1.columns))
    s2 = main_mod.strip_cols(pd.read_excel(xls, sheet_name=sheet2))

    docs = {main_mod.normalize_name(x) for x in list(s2.columns[1:])}

    date_col = s1.columns[0]
    s1[date_col] = pd.to_datetime(s1[date_col], errors="coerce").dt.normalize()

    hospital_cols = [c for c in s1.columns[1:] if str(c).strip() != "カテ当番"]

    slots = set()
    preassigned = {}
    dates = {}
    for row_pos, ridx in enumerate(s1.index):
        d = s1.at[ridx, date_col]
        if pd.isna(d):
            continue
        dates[row_pos] = pd.Timestamp(d).normalize()
        for hosp in hospital_cols:
            v = s1.at[ridx, hosp]
            v_norm = main_mod.normalize_name(v) if isinstance(v, str) else ""
            if v_norm in docs:
                slots.add((row_pos, hosp))
                preassigned[(row_pos, hosp)] = v_norm
            elif main_mod.is_slot_value(v):
                slots.add((row_pos, hosp))
    return docs, slots, preassigned, dates, hospital_cols


def _read_output_patterns(output_path, docs, slots, preassigned, dates):
    """出力xlsxの pattern_XX シートを独立に読み、枠ごとの割当を抽出する。

    レイアウト: 行1=バナー / 行2=ヘッダ / 行3.. = データ（sheet1と同じ行順）。
    """
    wb = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
    patterns = {}
    for sheet_name in wb.sheetnames:
        if not sheet_name.startswith("pattern_") or sheet_name.endswith("_summary"):
            continue
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        banner = rows[0][0] if rows else None
        header = list(rows[1]) if len(rows) > 1 else []
        col_index = {name: j for j, name in enumerate(header)}
        data_rows = rows[2:]  # 各要素が sheet1 の1行に対応

        assignments = {}  # (row_pos, colname) -> {"raw":..., "doc":..., "date":..., "pre":bool}
        for (row_pos, colname) in slots:
            j = col_index.get(colname)
            raw = None
            if j is not None and row_pos < len(data_rows):
                raw = data_rows[row_pos][j]
            doc = None
            if isinstance(raw, str):
                nn = raw.strip().replace(" ", "").replace("　", "")
                if nn in docs:
                    doc = nn
            assignments[(row_pos, colname)] = {
                "raw": raw,
                "doc": doc,
                "date": dates.get(row_pos),
                "pre": (row_pos, colname) in preassigned,
                "colname": colname,
            }
        patterns[sheet_name] = {
            "banner": banner,
            "header": header,
            "n_data_rows": len(data_rows),
            "assignments": assignments,
        }
    wb.close()
    return patterns


def _read_summary_grid(output_path, sheet_name):
    wb = openpyxl.load_workbook(output_path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


class Harness:
    """run() 実行結果と、独立に読み直した出力データをまとめて保持する。"""

    def __init__(self, main_mod, output_path, docs, slots, preassigned, dates, hospital_cols, patterns):
        self.main = main_mod
        self.output_path = output_path
        self.docs = docs
        self.slots = slots
        self.preassigned = preassigned
        self.dates = dates
        self.hospital_cols = hospital_cols
        self.patterns = patterns

    def pattern_sheets(self):
        return sorted(self.patterns.keys())

    def read_summary(self, sheet_name):
        return _read_summary_grid(self.output_path, sheet_name)


@pytest.fixture(scope="session")
def harness(tmp_path_factory):
    """main.run() をセッション中1回だけ実行し、出力を読み直した Harness を返す。"""
    import contextlib
    import io

    import main as main_mod

    assert os.path.exists(INPUT_XLSX), f"入力雛形が見つかりません: {INPUT_XLSX}"

    out_dir = tmp_path_factory.mktemp("tochoku_out")
    # run() は大量に print するのでキャプチャして握りつぶす（失敗時のみ再送出）。
    buf = io.StringIO()
    try:
        with contextlib.redirect_stdout(buf):
            output_path = main_mod.run(INPUT_XLSX, output_dir=str(out_dir), num_patterns=NUM_PATTERNS)
    except Exception:
        sys.stdout.write(buf.getvalue())
        raise

    assert os.path.exists(output_path), f"出力ファイルが生成されていません: {output_path}"

    docs, slots, preassigned, dates, hospital_cols = _read_input_slots(main_mod)
    patterns = _read_output_patterns(output_path, docs, slots, preassigned, dates)

    return Harness(main_mod, output_path, docs, slots, preassigned, dates, hospital_cols, patterns)
