"""v6.10.0 コア配分3件の回帰テスト（学年クォータ制 / ABS-012暦週 / ABS-016日直月1）。

docs/Tochoku.v7_2026.03.1.xlsx は変更禁止のため、目標列付きの合成入力を
openpyxl でその場生成して検証する。

合成入力の設計（2026-03-01(日)〜03-28(土) = 暦週ちょうど4週・医師6人・20枠）:
  - 医A〜医D（上級・属性2）: 大学目標1 + 外目標2 = 3回
  - 医E/医F（若手・属性1）: 大学目標2 + 外目標2 = 4回
  - 大学系(B-K)枠 8 = Σ大学目標 / 外病院(L-Y)枠 12 = Σ外目標
  - 日直枠（土曜昼/祝日昼/支援日直）3つ → ABS-016（月1回）が全員で成立可能
  - 大学系枠は各暦週2つ → ABS-012（暦週1回）で医師6人に分散可能
"""

import datetime
import os
import subprocess
import sys

import openpyxl
import pandas as pd
import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

import main as main_mod  # noqa: E402
import solver_cpsat  # noqa: E402

# 医師: (大学目標, 外目標)。若手（医E/医F）4回・上級（医A〜医D）3回の傾斜
TARGETS = {
    "医A": (1, 2), "医B": (1, 2), "医C": (1, 2), "医D": (1, 2),
    "医E": (2, 2), "医F": (2, 2),
}
ATTRS = {"医A": 2, "医B": 2, "医C": 2, "医D": 2, "医E": 1, "医F": 1}
DOCTORS = list(TARGETS)

SHEET1_COLS = [
    "Date", "大学平日", "大学土曜昼", "大学土曜夜", "大学日曜昼", "大学日曜夜",
    "大学祝日昼", "大学祝日夜", "支援平日", "支援日直", "支援当直",
    "病院X", "病院Y", "病院Z", "カテ当番",
]
UNIV_COLS = set(SHEET1_COLS[1:11])     # B-K相当
NICHOKU_COLS = {"大学土曜昼", "大学日曜昼", "大学祝日昼", "支援日直"}

# 日付 -> 枠を立てる列（2026-03-20(金)は祝日=春分の日）
SLOTS = {
    2: ["病院X"], 3: ["大学平日"], 4: ["病院Y"],
    7: ["大学土曜昼", "病院Z"], 8: ["大学日曜夜"], 9: ["病院X"],
    10: ["大学平日"], 11: ["病院Y"], 14: ["病院Z"],
    16: ["病院X"], 17: ["支援当直"], 18: ["病院Y"],
    20: ["大学祝日昼"], 21: ["病院Z"], 22: ["支援日直", "病院Z"],
    25: ["病院X"], 26: ["病院Y"], 28: ["大学土曜夜"],
}
N_SLOTS = sum(len(v) for v in SLOTS.values())  # 20
DATES = [datetime.datetime(2026, 3, d) for d in range(1, 29)]


def _sunday_week(ts):
    ts = pd.Timestamp(ts).normalize()
    return ts - pd.Timedelta(days=(ts.weekday() + 1) % 7)


def build_synthetic_xlsx(path, univ_col=True, ext_col=True,
                         blank_doc=None, univ_delta=None):
    """合成入力を生成する。

    Args:
        univ_col / ext_col: 目標列を出力するか（片列テスト用）
        blank_doc: この医師の「外目標」を空欄にする（空欄致命テスト用）
        univ_delta: {医師: 加算値} で大学目標をずらす（Σ不一致テスト用）
    """
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "sheet1"
    ws1.append(SHEET1_COLS)
    for dt in DATES:
        row = [dt] + [None] * (len(SHEET1_COLS) - 1)
        for col in SLOTS.get(dt.day, []):
            row[SHEET1_COLS.index(col)] = 1
        ws1.append(row)

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Date"] + DOCTORS)
    for dt in DATES:
        ws2.append([dt] + [1] * len(DOCTORS))

    ws3 = wb.create_sheet("Sheet3")
    header = ["氏名", "属性", "カテ当番", "出張日", "出張先",
              "全合計", "大学合計", "外病院合計", "平日", "休日合計"]
    if univ_col:
        header.append("大学目標")
    if ext_col:
        header.append("外目標")
    ws3.append(header)
    for doc in DOCTORS:
        univ, ext = TARGETS[doc]
        if univ_delta and doc in univ_delta:
            univ += univ_delta[doc]
        row = [doc, ATTRS[doc], None, None, None, 0, 0, 0, 0, 0]
        if univ_col:
            row.append(univ)
        if ext_col:
            row.append(None if blank_doc == doc else ext)
        ws3.append(row)

    wb.save(path)
    return path


# ---------------------------------------------------------------------------
# CP-SAT経路（solver_cpsat.solve）での傾斜配分・暦週・日直の機械検証
# ---------------------------------------------------------------------------

@pytest.fixture(scope="module")
def quota_solution(tmp_path_factory):
    path = str(tmp_path_factory.mktemp("quota") / "synthetic_quota.xlsx")
    build_synthetic_xlsx(path)
    data, solutions = solver_cpsat.solve(path, n_solutions=1, verbose=False)
    assert solutions, "CP-SATが解を返しませんでした"
    # 割当を (slot, doc) のリストに展開（固定割当なしの設計）
    assign = solutions[0]["assign"]
    picked = [(data.slots[si], doc) for si, doc in assign.items()]
    assert len(picked) == N_SLOTS
    return data, solutions[0], picked


def test_quota_enabled_and_targets_parsed(quota_solution):
    data, sol, _ = quota_solution
    assert data.quota_enabled is True
    assert data.univ_target == {d: t[0] for d, t in TARGETS.items()}
    assert data.ext_target == {d: t[1] for d, t in TARGETS.items()}
    assert data.extra_allowed == set()
    assert sol["status"] in ("OPTIMAL", "FEASIBLE")


def test_quota_tilted_distribution_realized(quota_solution):
    """若手4回・上級3回の傾斜配分＋大学系=大学目標が実際に出る（等式）。"""
    data, _, picked = quota_solution
    totals = {d: 0 for d in DOCTORS}
    univs = {d: 0 for d in DOCTORS}
    for slot, doc in picked:
        totals[doc] += 1
        if slot["hosp"] in UNIV_COLS:
            univs[doc] += 1
    for doc, (univ, ext) in TARGETS.items():
        assert totals[doc] == univ + ext, f"{doc}: 総数{totals[doc]} != 目標{univ + ext}"
        assert univs[doc] == univ, f"{doc}: 大学{univs[doc]} != 大学目標{univ}"
    # 傾斜（3回と4回の混在）が実在すること
    assert sorted(totals.values()) == [3, 3, 3, 3, 4, 4]


def test_quota_university_once_per_calendar_week(quota_solution):
    """ABS-012: 大学系は暦週(日曜始まり)で医師ごと1回まで。"""
    _, _, picked = quota_solution
    seen = set()  # (doc, week_start)
    for slot, doc in picked:
        if slot["hosp"] not in UNIV_COLS:
            continue
        key = (doc, _sunday_week(slot["date"]))
        assert key not in seen, f"暦週2回: {key}"
        seen.add(key)


def test_nichoku_at_most_once_per_month(quota_solution):
    """ABS-016: 日直（土曜昼/日曜昼/祝日昼/支援日直）は医師ごと月1回まで。"""
    _, _, picked = quota_solution
    counts = {}
    for slot, doc in picked:
        if slot["hosp"] in NICHOKU_COLS:
            counts[doc] = counts.get(doc, 0) + 1
    assert counts, "日直枠が1つも割り当てられていません（テスト空振り）"
    assert all(c <= 1 for c in counts.values()), f"日直2回以上: {counts}"


# ---------------------------------------------------------------------------
# 致命エラー（片列のみ / 一部空欄 / Σ不一致）
# ---------------------------------------------------------------------------

def test_single_target_column_is_fatal(tmp_path):
    path = str(tmp_path / "only_univ.xlsx")
    build_synthetic_xlsx(path, ext_col=False)
    with pytest.raises(ValueError, match="片方"):
        solver_cpsat.InputData(path, verbose=False)


def test_partially_blank_target_is_fatal(tmp_path):
    path = str(tmp_path / "blank.xlsx")
    build_synthetic_xlsx(path, blank_doc="医C")
    with pytest.raises(ValueError, match="医C.*空欄"):
        solver_cpsat.InputData(path, verbose=False)


def test_sum_mismatch_is_fatal_with_diff(tmp_path):
    path = str(tmp_path / "mismatch.xlsx")
    build_synthetic_xlsx(path, univ_delta={"医A": 1})  # Σ大学目標 9 != 大学枠 8
    with pytest.raises(ValueError, match=r"検算エラー.*差分 \+1"):
        solver_cpsat.InputData(path, verbose=False)


# ---------------------------------------------------------------------------
# main.py 側のパース関数（列が無い入力では従来動作 = enabled False）
# ---------------------------------------------------------------------------

def test_main_parse_quota_disabled_without_columns():
    df = pd.DataFrame({"氏名": DOCTORS, "全合計": [0] * 6})
    univ, ext, enabled = main_mod.parse_quota_targets(
        df, DOCTORS, {d: d for d in DOCTORS})
    assert enabled is False and univ == {} and ext == {}


def test_main_parse_quota_single_column_fatal():
    df = pd.DataFrame({"氏名": DOCTORS, "大学目標": [1] * 6})
    with pytest.raises(ValueError, match="片方"):
        main_mod.parse_quota_targets(df, DOCTORS, {d: d for d in DOCTORS})


def test_main_parse_quota_blank_fatal():
    df = pd.DataFrame({
        "氏名": DOCTORS,
        "大学目標": [1, 1, 1, 1, 2, 2],
        "外目標": [2, 2, None, 2, 2, 2],
    })
    with pytest.raises(ValueError, match="医C"):
        main_mod.parse_quota_targets(df, DOCTORS, {d: d for d in DOCTORS})


def test_main_validate_quota_totals_diff_display():
    univ = {d: t[0] for d, t in TARGETS.items()}
    ext = {d: t[1] for d, t in TARGETS.items()}
    # 一致 → 例外なし
    main_mod.validate_quota_totals(univ, ext, DOCTORS, 20, 8, 12)
    # 不一致 → 差分付き致命
    with pytest.raises(ValueError) as ei:
        main_mod.validate_quota_totals(univ, ext, DOCTORS, 21, 8, 12)
    assert "総枠数21と不一致" in str(ei.value) and "差分 -1" in str(ei.value)


# ---------------------------------------------------------------------------
# main.run() のエンドツーエンド（サブプロセス実行でグローバル汚染を回避）
# ---------------------------------------------------------------------------

def test_main_run_quota_end_to_end(tmp_path):
    input_path = str(tmp_path / "synthetic_quota.xlsx")
    build_synthetic_xlsx(input_path)
    out_dir = str(tmp_path / "out")
    code = (
        "import main; "
        f"p = main.run({input_path!r}, output_dir={out_dir!r}); "
        "print('OUTPUT::' + p)"
    )
    proc = subprocess.run(
        [sys.executable, "-c", code], cwd=REPO_ROOT,
        capture_output=True, text=True, timeout=300)
    assert proc.returncode == 0, f"main.run失敗:\n{proc.stdout[-3000:]}\n{proc.stderr[-3000:]}"
    assert "学年クォータ制" in proc.stdout
    assert "CP-SATで" in proc.stdout and "パターン生成" in proc.stdout

    out_path = next(line.split("OUTPUT::", 1)[1]
                    for line in proc.stdout.splitlines() if "OUTPUT::" in line)
    assert os.path.exists(out_path)

    wb = openpyxl.load_workbook(out_path, read_only=True, data_only=True)
    ws = wb["pattern_01"]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    banner, header = rows[0][0], list(rows[1])
    assert isinstance(banner, str) and banner.startswith("✅"), f"バナー: {banner!r}"

    col_index = {name: j for j, name in enumerate(header)}
    totals = {d: 0 for d in DOCTORS}
    univs = {d: 0 for d in DOCTORS}
    nichoku = {d: 0 for d in DOCTORS}
    for row in rows[2:]:
        for col in SHEET1_COLS[1:-1]:  # カテ当番以外の枠列
            v = row[col_index[col]]
            if isinstance(v, str) and v.strip() in TARGETS:
                doc = v.strip()
                totals[doc] += 1
                if col in UNIV_COLS:
                    univs[doc] += 1
                if col in NICHOKU_COLS:
                    nichoku[doc] += 1

    for doc, (univ, ext) in TARGETS.items():
        assert totals[doc] == univ + ext, f"{doc}: 総数{totals[doc]} != {univ + ext}"
        assert univs[doc] == univ, f"{doc}: 大学{univs[doc]} != {univ}"
    assert all(c <= 1 for c in nichoku.values()), f"日直2回以上: {nichoku}"
    assert sorted(totals.values()) == [3, 3, 3, 3, 4, 4]
