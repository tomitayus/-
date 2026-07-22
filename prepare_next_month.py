#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""prepare_next_month.py — 月次ワークフロー自動化ヘルパー（当直くん）

前月の当直表出力（pattern_NN シート）から医師別実績を集計し、当月雛形を
コピーして「翌月雛形ドラフト」を生成する。累計手転記による氏名照合事故
（最大の事故源）を排除するのが目的。

処理:
  (1) 出力xlsx の pattern_NN シートから医師別の当月実績を集計
      （全合計 / 大学合計(B-K) / 外病院合計(L-Y) / 平日 / 休日）
      分類ロジックは main.py の recompute_stats に一致させる。
  (2) 当月雛形をコピーし Sheet3 の累計列に加算
      （全合計・大学合計・外病院合計・平日・休日合計）
  (3) sheet1 の日付を翌月へ進め、曜日タイプ（月〜日 + 祝）ごとの枠数
      パターンを当月から推定して翌月グリッドを生成（祝日は自動判定）
  (4) Sheet2 のマークを全消去（翌月の希望は白紙）

氏名照合は normalize（空白除去）+ 完全一致のみ。曖昧一致はしない。
不一致医師は加算スキップして最後に警告一覧を表示する。

CLI:
  python3 prepare_next_month.py <当月出力.xlsx> <当月雛形.xlsx> \
      --pattern 1 -o <翌月雛形ドラフト.xlsx>

注意: 本スクリプトは main.py / tests/ の既存ファイルに一切依存しない
（独立実装）。分類定数のみ main.py の仕様を写経している。
"""

import argparse
import calendar
import collections
import datetime
import sys

import openpyxl

# ----------------------------------------------------------------------------
# ドラフト警告文（先頭シート & stdout に明示）
# ----------------------------------------------------------------------------
DRAFT_WARNING = (
    "⚠️ ドラフトです: 枠数パターンは推定。"
    "カテ当番(Z列)・出張変更・新入医師は手動確認"
)

# ----------------------------------------------------------------------------
# 内蔵祝日表（main.py と同一。jpholiday が無い環境でのフォールバック用）
# ----------------------------------------------------------------------------
_BUILTIN_HOLIDAYS = [
    # 2026年（内閣府発表の国民の祝日・休日）
    "2026-01-01", "2026-01-12", "2026-02-11", "2026-02-23", "2026-03-20",
    "2026-04-29", "2026-05-03", "2026-05-04", "2026-05-05", "2026-05-06",
    "2026-07-20", "2026-08-11", "2026-09-21", "2026-09-22", "2026-09-23",
    "2026-10-12", "2026-11-03", "2026-11-23",
    # 2027年
    "2027-01-01", "2027-01-11", "2027-02-11", "2027-02-23", "2027-03-21",
    "2027-03-22", "2027-04-29", "2027-05-03", "2027-05-04", "2027-05-05",
    "2027-07-19", "2027-08-11", "2027-09-20", "2027-09-23", "2027-10-11",
    "2027-11-03", "2027-11-23",
]

# 平日でも「休日枠」として扱う列インデックス（main.py recompute_stats と一致）
#   C_COL_INDEX=2, D_COL_INDEX=3, F_COL_INDEX=5, G_COL_INDEX=6
_HOLIDAY_LIKE_COL_IDX = frozenset({2, 3, 5, 6})

# 大学系（BG）= B〜K列 = 0-basedインデックス 1〜10
_BG_START, _BG_END = 1, 10
# 外病院（HT）= L〜Y列 = 0-basedインデックス 11〜24（上限）
_HT_START, _HT_END_CAP = 11, 24

_WD_JP = ["月", "火", "水", "木", "金", "土", "日"]

# 医師名として扱わない文字列（main.py SLOT_MARKERS の文字列系 + 未割当マーカー）
_NON_DOCTOR_STRINGS = frozenset({"1", "〇", "○", "◯", "◎", "UNASSIGNED"})


def normalize_name(name):
    """医師名を正規化（全角/半角スペース除去）。main.py と同一仕様。"""
    if name is None:
        return ""
    return str(name).strip().replace(" ", "").replace("　", "")


# ----------------------------------------------------------------------------
# 祝日判定
# ----------------------------------------------------------------------------
def build_holiday_set(year_months):
    """対象 (year, month) 集合に含まれる祝日の date 集合を返す。

    優先順: jpholiday（自動）→ 内蔵祝日表(2026-2027)。config.HOLIDAYS は
    常に追加マージ。main.py の祝日取得ロジックと同じ優先順位。
    """
    holidays = set()
    source = "内蔵祝日表(2026-2027)"
    try:
        import jpholiday  # type: ignore
        for (y, m) in sorted(year_months):
            for hd, _name in jpholiday.month_holidays(y, m):
                holidays.add(hd)
        source = "jpholiday（自動）"
    except ImportError:
        for hs in _BUILTIN_HOLIDAYS:
            d = datetime.date(*map(int, hs.split("-")))
            if (d.year, d.month) in year_months:
                holidays.add(d)
        if any(y > 2027 for (y, m) in year_months):
            print(
                "⚠️ WARNING: 2028年以降の祝日は内蔵表にありません。"
                "`pip install jpholiday` を検討してください"
            )
    # config.py の祝日を追加マージ（手動追加・上書き用。main.py と同一運用）
    try:
        import config as _cfg  # type: ignore
        for hs in getattr(_cfg, "HOLIDAYS", []) or []:
            d = datetime.date(*map(int, str(hs).split("-")))
            if (d.year, d.month) in year_months:
                holidays.add(d)
    except (ImportError, ValueError, TypeError):
        pass
    return holidays, source


def is_holiday(d, holiday_set):
    if isinstance(d, datetime.datetime):
        d = d.date()
    return d in holiday_set


def day_key(d, holiday_set):
    """曜日タイプキー: 祝日→'祝'、それ以外→曜日（月〜日）。"""
    if is_holiday(d, holiday_set):
        return "祝"
    if isinstance(d, datetime.datetime):
        d = d.date()
    return _WD_JP[d.weekday()]


# ----------------------------------------------------------------------------
# (1) 出力 pattern_NN シートから医師別の当月実績を集計
# ----------------------------------------------------------------------------
def _find_header_row(ws):
    """ヘッダー行（col A が 'Date'/'日付'）を探す。バナー行(1行目)対策。"""
    for r in range(1, min(ws.max_row, 8) + 1):
        v = ws.cell(row=r, column=1).value
        if isinstance(v, str) and v.strip().lower() in ("date", "日付"):
            return r
    # フォールバック: 2行目（バナー+ヘッダーの既定レイアウト）
    return 2


def _parse_grid_date(v):
    """pattern シートの日付セルを date に変換。

    出力は 'YYYY/M/D (曜日)' 文字列。datetime の場合も許容。
    パースできなければ None。
    """
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    s = str(v).strip()
    if not s:
        return None
    head = s.split(" ")[0].split("(")[0].strip()  # "2026/3/1"
    for sep in ("/", "-"):
        if sep in head:
            parts = head.split(sep)
            if len(parts) == 3:
                try:
                    return datetime.date(int(parts[0]), int(parts[1]), int(parts[2]))
                except ValueError:
                    return None
    return None


def collect_month_stats(output_path, pattern_num, holiday_set):
    """pattern_NN シートを読み、正規化名 → {total,bg,ht,wd,we} の dict を返す。

    分類は main.py recompute_stats と一致:
      - BG(大学)   : 列index 1..10
      - HT(外病院) : 列index 11..min(24, カテ当番手前)
      - 休日フラグ : is_holiday(date) or 土日 or (平日 and 列index∈{2,3,5,6})
    """
    wb = openpyxl.load_workbook(output_path, data_only=True)
    sheet_name = f"pattern_{pattern_num:02d}"
    if sheet_name not in wb.sheetnames:
        # 大小/桁揺れの救済（pattern_1 等）
        alt = f"pattern_{pattern_num}"
        cand = [s for s in wb.sheetnames
                if s.lower() == sheet_name.lower() or s.lower() == alt.lower()]
        if not cand:
            raise ValueError(
                f"❌ 出力に '{sheet_name}' シートがありません。"
                f"存在するシート: {wb.sheetnames}"
            )
        sheet_name = cand[0]
    ws = wb[sheet_name]

    header_row = _find_header_row(ws)
    headers = [ws.cell(row=header_row, column=c).value
               for c in range(1, ws.max_column + 1)]  # 1-based → list 0-based

    # カテ当番列（0-based grid index）を特定 → 病院列の右端を決定
    kate_j = None
    for j, h in enumerate(headers):
        if isinstance(h, str) and h.strip() == "カテ当番":
            kate_j = j
            break
    if kate_j is not None:
        last_hosp_j = kate_j - 1
    else:
        # カテ当番が無い場合: 末尾の非空ヘッダーまで
        last_hosp_j = 0
        for j, h in enumerate(headers):
            if h is not None and str(h).strip() != "" and not str(h).startswith("【"):
                last_hosp_j = j
    ht_end = min(_HT_END_CAP, last_hosp_j)

    # 集計対象の列 index（1..last_hosp_j、カテ当番は除外）
    stats = collections.defaultdict(
        lambda: {"total": 0, "bg": 0, "ht": 0, "wd": 0, "we": 0}
    )
    seen_names = set()

    for r in range(header_row + 1, ws.max_row + 1):
        d = _parse_grid_date(ws.cell(row=r, column=1).value)
        if d is None:
            continue
        dow = d.weekday()
        for j in range(1, last_hosp_j + 1):
            if kate_j is not None and j == kate_j:
                continue
            val = ws.cell(row=r, column=j + 1).value  # openpyxl は 1-based
            if not isinstance(val, str):
                continue
            doc = normalize_name(val)
            if not doc or doc in _NON_DOCTOR_STRINGS:
                continue  # 未割当マーカー等は医師名ではない
            seen_names.add(doc)
            s = stats[doc]
            s["total"] += 1
            if _BG_START <= j <= _BG_END:
                s["bg"] += 1
            elif _HT_START <= j <= ht_end:
                s["ht"] += 1
            holi = (
                is_holiday(d, holiday_set)
                or dow >= 5
                or (dow < 5 and j in _HOLIDAY_LIKE_COL_IDX)
            )
            if holi:
                s["we"] += 1
            else:
                s["wd"] += 1

    return dict(stats), seen_names


# ----------------------------------------------------------------------------
# (3) 枠数パターン推定（当月グリッドから）
# ----------------------------------------------------------------------------
def _mode_ties_max(values):
    """最頻値（同数タイは大きい方）。空なら 0。"""
    if not values:
        return 0
    cnt = collections.Counter(values)
    best_freq = max(cnt.values())
    return max(v for v, f in cnt.items() if f == best_freq)


def estimate_frame_patterns(ws_sheet1, date_col, headers, kate_j, cur_holidays):
    """当月 sheet1 から (列index, 曜日キー) → 推定枠数 を返す。

    値は最頻値（タイは大きい方）。カテ当番列は推定対象外。
    """
    # 列ごとに、曜日キー別の枠数リストを収集
    by_col_key = collections.defaultdict(lambda: collections.defaultdict(list))
    for r in range(2, ws_sheet1.max_row + 1):
        d = ws_sheet1.cell(row=r, column=1).value
        if not isinstance(d, (datetime.datetime, datetime.date)):
            continue
        dd = d.date() if isinstance(d, datetime.datetime) else d
        dk = day_key(dd, cur_holidays)
        for j, h in enumerate(headers):
            if j == 0:
                continue  # 日付列
            if kate_j is not None and j == kate_j:
                continue
            v = ws_sheet1.cell(row=r, column=j + 1).value
            if isinstance(v, (int, float)):
                v = int(v)
            elif isinstance(v, str) and v.strip() != "":
                # 医師名直書き（固定割当）や 〇 マーカーは 1枠として数える
                v = 1
            else:
                v = 0
            by_col_key[j][dk].append(v)

    est = {}
    for j, keymap in by_col_key.items():
        est[j] = {k: _mode_ties_max(vals) for k, vals in keymap.items()}
    has_holiday_sample = any("祝" in km for km in est.values())
    return est, has_holiday_sample


def _frame_value(est, j, dk, headers, has_holiday_sample):
    """曜日キー dk における列 j の推定枠数。

    '祝' キーの標本が当月に無い場合のフォールバック:
      - 大学祝日昼/夜 → 大学日曜昼/夜 の日曜推定を流用
      - 大学平日/大学土曜/大学日曜 → 0
      - その他 → 日曜推定
    """
    coljmap = est.get(j, {})
    if dk in coljmap:
        return coljmap[dk]
    if dk == "祝" and not has_holiday_sample:
        name = headers[j] if j < len(headers) else None
        name = str(name).strip() if name is not None else ""
        # 名前→列index 逆引き
        name_to_j = {str(h).strip(): i for i, h in enumerate(headers)
                     if isinstance(h, str)}
        if name == "大学祝日昼":
            src = name_to_j.get("大学日曜昼")
            return est.get(src, {}).get("日", 0) if src is not None else 0
        if name == "大学祝日夜":
            src = name_to_j.get("大学日曜夜")
            return est.get(src, {}).get("日", 0) if src is not None else 0
        if name in ("大学平日", "大学土曜昼", "大学土曜夜",
                    "大学日曜昼", "大学日曜夜"):
            return 0
        return coljmap.get("日", 0)
    return 0


# ----------------------------------------------------------------------------
# メイン処理
# ----------------------------------------------------------------------------
def _month_range_dates(year, month):
    ndays = calendar.monthrange(year, month)[1]
    return [datetime.date(year, month, d) for d in range(1, ndays + 1)]


def _next_month(year, month):
    return (year + 1, 1) if month == 12 else (year, month + 1)


def _detect_current_month(ws_sheet1):
    dates = []
    for r in range(2, ws_sheet1.max_row + 1):
        v = ws_sheet1.cell(row=r, column=1).value
        if isinstance(v, (datetime.datetime, datetime.date)):
            dd = v.date() if isinstance(v, datetime.datetime) else v
            dates.append(dd)
    if not dates:
        raise ValueError("❌ 当月雛形 sheet1 に日付が見つかりません")
    first = min(dates)
    return first.year, first.month, dates


def _find_sheet(wb, target):
    """名前の大小・表記ゆれに耐えるシート取得。"""
    if target in wb.sheetnames:
        return wb[target]
    low = {s.lower(): s for s in wb.sheetnames}
    if target.lower() in low:
        return wb[low[target.lower()]]
    return None


def prepare(output_path, template_path, pattern_num, out_path):
    # --- 当月雛形を読み込み（書式保持のため values_only にしない）---
    wb = openpyxl.load_workbook(template_path)
    ws1 = _find_sheet(wb, "sheet1")
    ws2 = _find_sheet(wb, "Sheet2")
    # 医師情報シート: main.py と同じ優先順（sheet4 があればそちら、無ければ Sheet3）
    ws3 = _find_sheet(wb, "sheet4") or _find_sheet(wb, "Sheet3")
    if ws1 is None or ws2 is None or ws3 is None:
        raise ValueError(
            f"❌ 雛形に sheet1/Sheet2/Sheet3 が揃っていません: {wb.sheetnames}"
        )

    cur_year, cur_month, cur_dates = _detect_current_month(ws1)
    nxt_year, nxt_month = _next_month(cur_year, cur_month)

    # --- 祝日集合（当月＋翌月）---
    cur_ym = {(cur_year, cur_month)}
    nxt_ym = {(nxt_year, nxt_month)}
    cur_holidays, _src_c = build_holiday_set(cur_ym)
    nxt_holidays, src_n = build_holiday_set(nxt_ym)

    # --- (1) 当月実績集計（休日分類には当月の祝日集合を使う）---
    month_stats, seen_names = collect_month_stats(
        output_path, pattern_num, cur_holidays
    )

    # --- ヘッダー & カテ当番列の把握（sheet1）---
    headers = [ws1.cell(row=1, column=c).value for c in range(1, ws1.max_column + 1)]
    kate_j = None
    for j, h in enumerate(headers):
        if isinstance(h, str) and h.strip() == "カテ当番":
            kate_j = j
            break

    # --- 名簿（Sheet3 氏名列）---
    roster = []
    roster_norm = {}
    for r in range(2, ws3.max_row + 1):
        nm = ws3.cell(row=r, column=1).value
        if nm is None or str(nm).strip() == "":
            continue
        nn = normalize_name(nm)
        roster.append((r, nm, nn))
        roster_norm[nn] = r

    # --- (2) Sheet3 累計加算（完全一致のみ）---
    # 累計列の見出し → 列番号（1-based）
    hdr3 = {ws3.cell(row=1, column=c).value: c
            for c in range(1, ws3.max_column + 1)}
    cum_map = {
        "全合計": "total",
        "大学合計": "bg",
        "外病院合計": "ht",
        "平日": "wd",
        "休日合計": "we",
    }
    missing_cols = [k for k in cum_map if k not in hdr3]
    if missing_cols:
        raise ValueError(f"❌ Sheet3 に累計列が見つかりません: {missing_cols}")

    audit = []  # 検算用: (氏名, 列, 旧累計, 当月実績, 新累計)
    for (r, nm, nn) in roster:
        st = month_stats.get(nn)
        if st is None:
            st = {"total": 0, "bg": 0, "ht": 0, "wd": 0, "we": 0}
        for col_label, stat_key in cum_map.items():
            c = hdr3[col_label]
            old = ws3.cell(row=r, column=c).value
            old_num = float(old) if isinstance(old, (int, float)) else 0.0
            add = st[stat_key]
            new = old_num + add
            # 整数なら int で書き戻す（見た目維持）
            new_val = int(new) if float(new).is_integer() else new
            ws3.cell(row=r, column=c, value=new_val)
            audit.append((nm, col_label, old_num, add, new_val))

    # --- 照合不一致（出力に居るが名簿に無い医師）---
    unmatched = sorted(n for n in seen_names if n not in roster_norm)

    # --- (3) sheet1 グリッドを翌月へ ---
    date_num_fmt = ws1.cell(row=2, column=1).number_format  # 日付書式を継承
    est, has_holiday_sample = estimate_frame_patterns(
        ws1, 1, headers, kate_j, cur_holidays
    )
    nxt_dates = _month_range_dates(nxt_year, nxt_month)

    # 既存データ行のクリア範囲（当月と翌月の長い方）
    # 注意: openpyxl の ws.cell(..., value=None) は no-op のため .value 直接代入で消す
    n_clear = max(len(cur_dates), len(nxt_dates))
    for i in range(n_clear):
        r = 2 + i
        for j in range(len(headers)):
            ws1.cell(row=r, column=j + 1).value = None

    # 翌月グリッド書き込み
    nxt_holiday_hits = []
    for i, d in enumerate(nxt_dates):
        r = 2 + i
        dk = day_key(d, nxt_holidays)
        if dk == "祝":
            nxt_holiday_hits.append(d)
        cell = ws1.cell(row=r, column=1,
                        value=datetime.datetime(d.year, d.month, d.day))
        cell.number_format = date_num_fmt
        for j, h in enumerate(headers):
            if j == 0:
                continue
            if kate_j is not None and j == kate_j:
                # カテ当番は手動記入（空欄）
                ws1.cell(row=r, column=j + 1).value = None
                continue
            val = _frame_value(est, j, dk, headers, has_holiday_sample)
            ws1.cell(row=r, column=j + 1, value=int(val))

    # --- (4) Sheet2 マーク全消去 + 日付を翌月へ ---
    s2_date_fmt = ws2.cell(row=2, column=1).number_format
    n2_clear = 0
    for r in range(2, ws2.max_row + 1):
        v = ws2.cell(row=r, column=1).value
        if isinstance(v, (datetime.datetime, datetime.date)) or n2_clear < n_clear:
            n2_clear += 1
        else:
            break
    n2_clear = max(n2_clear, len(cur_dates))
    # まず全データ行をクリア（value=None 引数渡しは no-op のため直接代入）
    for i in range(max(n2_clear, len(nxt_dates))):
        r = 2 + i
        for c in range(1, ws2.max_column + 1):
            ws2.cell(row=r, column=c).value = None
    # 翌月の日付だけ書き戻す（医師列は白紙）
    for i, d in enumerate(nxt_dates):
        r = 2 + i
        cell = ws2.cell(row=r, column=1,
                        value=datetime.datetime(d.year, d.month, d.day))
        cell.number_format = s2_date_fmt

    # --- 先頭シートにドラフト警告 ---
    readme = wb.create_sheet("⚠️README", 0)
    readme["A1"] = DRAFT_WARNING
    readme["A3"] = f"当月: {cur_year}/{cur_month}  →  翌月: {nxt_year}/{nxt_month}"
    readme["A4"] = f"元出力: pattern_{pattern_num:02d}"
    readme["A5"] = f"翌月祝日({src_n}): " + ", ".join(
        d.strftime("%m/%d") for d in sorted(nxt_holidays)
    )
    if unmatched:
        readme["A7"] = "⚠️ 名簿未照合（累計加算スキップ）: " + ", ".join(unmatched)

    wb.save(out_path)

    return {
        "cur_ym": (cur_year, cur_month),
        "nxt_ym": (nxt_year, nxt_month),
        "audit": audit,
        "unmatched": unmatched,
        "nxt_holidays": sorted(nxt_holidays),
        "nxt_holiday_hits": sorted(set(nxt_holiday_hits)),
        "src_holiday": src_n,
        "roster_size": len(roster),
        "month_stats": month_stats,
        "out_path": out_path,
    }


def main(argv=None):
    ap = argparse.ArgumentParser(
        description="前月出力→翌月雛形ドラフトを生成（累計自動転記）"
    )
    ap.add_argument("output_xlsx", help="当月の当直表出力 xlsx（pattern_NN 入り）")
    ap.add_argument("template_xlsx", help="当月の雛形 xlsx")
    ap.add_argument("--pattern", type=int, default=1,
                    help="使用する pattern 番号（既定 1）")
    ap.add_argument("-o", "--output", required=True, help="翌月雛形ドラフトの出力先")
    args = ap.parse_args(argv)

    print("=" * 60)
    print("  当直くん 月次準備ヘルパー: prepare_next_month")
    print("=" * 60)
    print(DRAFT_WARNING)
    print("-" * 60)

    res = prepare(args.output_xlsx, args.template_xlsx, args.pattern, args.output)

    cy, cm = res["cur_ym"]
    ny, nm = res["nxt_ym"]
    print(f"当月 {cy}/{cm} → 翌月 {ny}/{nm}")
    print(f"名簿医師数: {res['roster_size']}")
    print(f"翌月祝日（{res['src_holiday']}）: " +
          ", ".join(d.strftime('%m/%d') for d in res['nxt_holidays']))
    if res["unmatched"]:
        print("⚠️ 名簿未照合（累計加算スキップ）: " + ", ".join(res["unmatched"]))
    else:
        print("✅ 出力医師は全員名簿と完全一致（累計加算済み）")
    print(f"出力: {res['out_path']}")
    print("=" * 60)
    return 0


if __name__ == "__main__":
    sys.exit(main())
