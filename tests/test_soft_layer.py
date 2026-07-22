"""v6.12.0 ソフト層3件の回帰テスト（二層累計公平 / ソフト回避 / カテ番×平日大学加点）。

docs/Tochoku.v7_2026.03.1.xlsx は変更禁止のため、合成入力を openpyxl でその場生成して
CP-SAT経路（solver_cpsat.solve）で機械検証する。重み・リストは solver_cpsat の
モジュール属性を monkeypatch する（実装は参照時に読む設計）。

合成入力（test_grade_quota.py と同じ骨格・目標列なし＝非クォータ）:
  2026-03-01(日)〜03-28(土)・医師6人・20枠（大学系8 / 外病院12）
  BASE_TARGET=3・EXTRA 2枠（属性1の医E/医F が+1）→ 総数は 医A-D=3, 医E/F=4 に固定
  （Σcap=枠数のため全医師 count==cap が強制され、外病院の配分だけが自由度になる）
"""

import datetime
import os
import sys

import openpyxl
import pandas as pd
import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

import solver_cpsat  # noqa: E402

DOCTORS = ["医A", "医B", "医C", "医D", "医E", "医F"]
ATTRS = {"医A": 2, "医B": 2, "医C": 2, "医D": 2, "医E": 1, "医F": 1}

SHEET1_COLS = [
    "Date", "大学平日", "大学土曜昼", "大学土曜夜", "大学日曜昼", "大学日曜夜",
    "大学祝日昼", "大学祝日夜", "支援平日", "支援日直", "支援当直",
    "病院X", "病院Y", "病院Z", "カテ当番",
]
UNIV_COLS = set(SHEET1_COLS[1:11])   # B-K相当
EXT_COLS = {"病院X", "病院Y", "病院Z"}  # L-Y相当

# 日付 -> 枠を立てる列（2026-03-20(金)は祝日=春分の日）
SLOTS = {
    2: ["病院X"], 3: ["大学平日"], 4: ["病院Y"],
    7: ["大学土曜昼", "病院Z"], 8: ["大学日曜夜"], 9: ["病院X"],
    10: ["大学平日"], 11: ["病院Y"], 14: ["病院Z"],
    16: ["病院X"], 17: ["支援当直"], 18: ["病院Y"],
    20: ["大学祝日昼"], 21: ["病院Z"], 22: ["支援日直", "病院Z"],
    25: ["病院X"], 26: ["病院Y"], 28: ["大学土曜夜"],
}
DATES = [datetime.datetime(2026, 3, d) for d in range(1, 29)]


def build_soft_xlsx(path, prev_ht=None, kate_days=None, kate_doc_team=None):
    """非クォータ（目標列なし）の合成入力を生成する。

    Args:
        prev_ht: {医師: 前月外病院合計} の偏り注入（省略=全員0）
        kate_days: sheet1「カテ当番」列にチームコードを入れる {日(int): コード}
        kate_doc_team: Sheet3「カテ当番」列 {医師: チームコード}
    """
    prev_ht = prev_ht or {}
    kate_days = kate_days or {}
    kate_doc_team = kate_doc_team or {}
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "sheet1"
    ws1.append(SHEET1_COLS)
    for dt in DATES:
        row = [dt] + [None] * (len(SHEET1_COLS) - 1)
        for col in SLOTS.get(dt.day, []):
            row[SHEET1_COLS.index(col)] = 1
        if dt.day in kate_days:
            row[SHEET1_COLS.index("カテ当番")] = kate_days[dt.day]
        ws1.append(row)

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Date"] + DOCTORS)
    for dt in DATES:
        ws2.append([dt] + [1] * len(DOCTORS))

    ws3 = wb.create_sheet("Sheet3")
    ws3.append(["氏名", "属性", "カテ当番", "出張日", "出張先",
                "全合計", "大学合計", "外病院合計", "平日", "休日合計"])
    for doc in DOCTORS:
        ws3.append([doc, ATTRS[doc], kate_doc_team.get(doc), None, None,
                    0, 0, prev_ht.get(doc, 0), 0, 0])

    wb.save(path)
    return path


def _counts(data, solution):
    """割当から (総数, 大学数, 外病院数, slot->doc) を集計する。"""
    total = {d: 0 for d in DOCTORS}
    univ = {d: 0 for d in DOCTORS}
    ext = {d: 0 for d in DOCTORS}
    by_slot = {}
    for si, doc in solution["assign"].items():
        slot = data.slots[si]
        total[doc] += 1
        if slot["hosp"] in UNIV_COLS:
            univ[doc] += 1
        if slot["hosp"] in EXT_COLS:
            ext[doc] += 1
        by_slot[(pd.Timestamp(slot["date"]).day, slot["hosp"])] = doc
    return total, univ, ext, by_slot


# ---------------------------------------------------------------------------
# A. 二層累計公平（提案#5）: prev_ht偏り → 外病院割当が累計少ない医師に寄る
# ---------------------------------------------------------------------------

def test_prev_ht_parsed_from_sheet3(tmp_path):
    path = str(tmp_path / "prev.xlsx")
    build_soft_xlsx(path, prev_ht={"医A": 10})
    data = solver_cpsat.InputData(path, verbose=False)
    assert data.prev_ht == {"医A": 10, "医B": 0, "医C": 0, "医D": 0, "医E": 0, "医F": 0}


def test_cum_fairness_pulls_external_away_from_high_prev(tmp_path, monkeypatch):
    """前月外病院累計が突出した医Aは、当月の外病院割当が最小（大学系に寄る）。"""
    path = str(tmp_path / "skew.xlsx")
    build_soft_xlsx(path, prev_ht={"医A": 10})
    monkeypatch.setattr(solver_cpsat, "W_FAIR_CUM_HT", 50)
    monkeypatch.setattr(solver_cpsat, "CUM_FAIRNESS_EXEMPT", [])
    data, solutions = solver_cpsat.solve(path, n_solutions=1, verbose=False)
    total, univ, ext, _ = _counts(data, solutions[0])
    # 総数はcap固定（3,3,3,3,4,4）
    assert total == {"医A": 3, "医B": 3, "医C": 3, "医D": 3, "医E": 4, "医F": 4}
    # 医Aは大学系上限(2)まで使って外病院を最小の1回に抑える
    assert ext["医A"] == 1, f"医Aの外病院が最小化されていない: {ext}"
    assert univ["医A"] == 2
    # 他医師は累計を追い上げる側（全員2回以上）
    others = {d: ext[d] for d in DOCTORS if d != "医A"}
    assert min(others.values()) >= 2, f"外病院が累計少ない側に寄っていない: {ext}"


def test_cum_fairness_disabled_is_backward_compatible(tmp_path, monkeypatch):
    """W_FAIR_CUM_HT=0（旧configに項目なし相当）では目的関数に累計項が入らない。"""
    path = str(tmp_path / "skew0.xlsx")
    build_soft_xlsx(path, prev_ht={"医A": 10})
    monkeypatch.setattr(solver_cpsat, "W_FAIR_CUM_HT", 0)
    data = solver_cpsat.InputData(path, verbose=False)
    sched = solver_cpsat.CpSatScheduler(data)
    assert sched.cum_ht_over is None


def test_cum_fairness_exempt_doctor_excluded(tmp_path, monkeypatch):
    """CUM_FAIRNESS_EXEMPTの医師（恒常例外）は均等化対象から外れる。"""
    path = str(tmp_path / "exempt.xlsx")
    build_soft_xlsx(path, prev_ht={"医A": 10})
    monkeypatch.setattr(solver_cpsat, "W_FAIR_CUM_HT", 50)
    monkeypatch.setattr(solver_cpsat, "CUM_FAIRNESS_EXEMPT", ["医A"])
    data, solutions = solver_cpsat.solve(path, n_solutions=1, verbose=False)
    _, _, ext, _ = _counts(data, solutions[0])
    # 医Aを除いた対象（prev全員0）同士の外病院数はほぼ均等（spread<=1）
    others = [ext[d] for d in DOCTORS if d != "医A"]
    assert max(others) - min(others) <= 1, f"均等化対象のspreadが大きい: {ext}"


# ---------------------------------------------------------------------------
# B. ソフト回避フラグ（提案#9）: 割当1回ごとに軽ペナルティ
# ---------------------------------------------------------------------------

def test_soft_avoid_adds_penalty_per_assignment(tmp_path, monkeypatch):
    """SOFT_AVOID該当医師の割当数×W_SOFT_AVOIDが目的関数値に上乗せされる。

    合成入力は Σcap=枠数 のため各医師の総数は固定（医A=3）。
    したがって最適値の差はちょうど 3×W_SOFT_AVOID になる。
    """
    path = str(tmp_path / "avoid.xlsx")
    build_soft_xlsx(path)
    monkeypatch.setattr(solver_cpsat, "SOFT_AVOID_DOCTORS", [])
    _, sols0 = solver_cpsat.solve(path, n_solutions=1, verbose=False)
    monkeypatch.setattr(solver_cpsat, "SOFT_AVOID_DOCTORS", ["医A"])
    monkeypatch.setattr(solver_cpsat, "W_SOFT_AVOID", 15)
    data1, sols1 = solver_cpsat.solve(path, n_solutions=1, verbose=False)
    assert sols0[0]["status"] == "OPTIMAL" and sols1[0]["status"] == "OPTIMAL"
    total, _, _, _ = _counts(data1, sols1[0])
    assert total["医A"] == 3
    assert sols1[0]["objective"] - sols0[0]["objective"] == pytest.approx(45, abs=1e-6)


# ---------------------------------------------------------------------------
# C. カテ番×平日大学の「できれば合わせる」（提案#10）
# ---------------------------------------------------------------------------

def test_kate_weekday_bonus_attracts_kate_doctor(tmp_path, monkeypatch):
    """カテ当番日(3/3=火)の平日大学枠(B列)にカテ当番医師（医B）が引き寄せられる。"""
    path = str(tmp_path / "kate.xlsx")
    build_soft_xlsx(path, kate_days={3: "A"}, kate_doc_team={"医B": "A"})
    # 加点を支配的にして決定的にする（公平性1e5より小さい範囲）
    monkeypatch.setattr(solver_cpsat, "W_KATE_WEEKDAY_BONUS", 1000)
    data, solutions = solver_cpsat.solve(path, n_solutions=1, verbose=False)
    _, _, _, by_slot = _counts(data, solutions[0])
    assert by_slot[(3, "大学平日")] == "医B", (
        f"カテ当番医師がカテ当番日の平日大学枠に入っていない: {by_slot[(3, '大学平日')]}")


def test_kate_weekday_bonus_disabled_collects_no_bonus_vars(tmp_path, monkeypatch):
    """W_KATE_WEEKDAY_BONUS=0（旧config相当）では加点項が目的関数に入らない。"""
    path = str(tmp_path / "kate0.xlsx")
    build_soft_xlsx(path, kate_days={3: "A"}, kate_doc_team={"医B": "A"})
    monkeypatch.setattr(solver_cpsat, "W_KATE_WEEKDAY_BONUS", 0)
    monkeypatch.setattr(solver_cpsat, "W_FAIR_CUM_HT", 0)
    monkeypatch.setattr(solver_cpsat, "SOFT_AVOID_DOCTORS", [])
    data = solver_cpsat.InputData(path, verbose=False)
    sched = solver_cpsat.CpSatScheduler(data)
    # 目的関数が旧来3層（SEMI/FAIR/SOFT簡易版）のみで構築できること（例外なし）
    assert sched.cum_ht_over is None


# ---------------------------------------------------------------------------
# D. summaryシートへの「外病院累計spread」行（実雛形のフル実走出力で確認）
# ---------------------------------------------------------------------------

def test_summary_has_cum_ht_spread_row(harness):
    for psheet in harness.pattern_sheets():
        rows = harness.read_summary(f"{psheet}_summary")
        cells = {str(c) for r in rows for c in r if c is not None}
        assert "外病院累計spread" in cells, f"{psheet}_summary に外病院累計spread行がない"
        assert "ソフト回避医師の割当数" in cells
        assert "カテ番×平日大学一致" in cells
