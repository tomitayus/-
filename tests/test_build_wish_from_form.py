# -*- coding: utf-8 -*-
"""build_wish_from_form.py のテスト（Form回答→入力.xlsx 反映）。

合成の入力xlsx（sheet1/Sheet2/Sheet3）と回答CSVを tmp_path 上に作り、
Sheet2 の 0・「希望」シートの ×/○・競合(不可優先)・氏名不一致・未回答を検証する。
"""

import datetime
import os
import sys

import openpyxl
import pandas as pd
import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

import build_wish_from_form as bwf

DOCTORS = ["医A", "医B", "医C"]
DATES = [datetime.date(2026, 3, d) for d in range(1, 29)]


def build_input(path, with_wish=False):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "sheet1"
    ws1.append(["Date", "病院X"])
    for d in DATES:
        ws1.append([datetime.datetime(d.year, d.month, d.day), 1])
    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Date"] + DOCTORS)
    for d in DATES:
        ws2.append([datetime.datetime(d.year, d.month, d.day)] + [None] * len(DOCTORS))
    ws3 = wb.create_sheet("Sheet3")
    ws3.append(["氏名"])
    for doc in DOCTORS:
        ws3.append([doc])
    if with_wish:
        wsw = wb.create_sheet("希望")
        wsw.append(["Date"] + DOCTORS)
        for d in DATES:
            wsw.append([datetime.datetime(d.year, d.month, d.day)] + [None] * len(DOCTORS))
    wb.save(path)
    return str(path)


def write_responses(path):
    rows = [
        {"タイムスタンプ": "2026/03/01 9:00", "お名前": "医A", "対象月": "2026年3月",
         "当直できない日": "3/2(月), 3/9(月)",
         "できれば避けたい日": "3/2(月), 3/16(月)",   # 3/2 は不可と競合
         "できれば入りたい日": "3/23(月)"},
        {"タイムスタンプ": "2026/03/01 9:05", "お名前": "医B", "対象月": "2026年3月",
         "当直できない日": "",
         "できれば避けたい日": "3/3(火)",
         "できれば入りたい日": ""},
        {"タイムスタンプ": "2026/03/01 9:10", "お名前": "医不在", "対象月": "2026年3月",
         "当直できない日": "3/5(木)", "できれば避けたい日": "", "できれば入りたい日": ""},
    ]
    pd.DataFrame(rows).to_csv(path, index=False)
    return str(path)


# ---------------------------------------------------------------------------
# 純粋関数
# ---------------------------------------------------------------------------
def test_parse_labels():
    assert bwf.parse_labels("3/2(月), 3/9(月)") == ["3/2(月)", "3/9(月)"]
    assert bwf.parse_labels("3/2、3/3\n3/4") == ["3/2", "3/3", "3/4"]
    assert bwf.parse_labels("") == []
    assert bwf.parse_labels(None) == []
    assert bwf.parse_labels(float("nan")) == []


def test_label_to_date():
    dates = DATES
    by_ymd = {(d.year, d.month, d.day): d for d in dates}
    by_md = {(d.month, d.day): d for d in dates}
    by_d = {d.day: d for d in dates}
    assert bwf.label_to_date("3/2(月)", by_ymd, by_md, by_d) == datetime.date(2026, 3, 2)
    assert bwf.label_to_date("2026/3/2", by_ymd, by_md, by_d) == datetime.date(2026, 3, 2)
    assert bwf.label_to_date("2", by_ymd, by_md, by_d) == datetime.date(2026, 3, 2)
    assert bwf.label_to_date("なし", by_ymd, by_md, by_d) is None


def test_find_col():
    cols = ["タイムスタンプ", "お名前", "当直できない日", "できれば避けたい日", "できれば入りたい日"]
    assert bwf.find_col(cols, bwf.NAME_KEYS) == "お名前"
    assert bwf.find_col(cols, bwf.CANT_KEYS) == "当直できない日"
    assert bwf.find_col(cols, bwf.AVOID_KEYS) == "できれば避けたい日"
    assert bwf.find_col(cols, bwf.WANT_KEYS) == "できれば入りたい日"


# ---------------------------------------------------------------------------
# エンドツーエンド
# ---------------------------------------------------------------------------
def _read(out_path):
    wb = openpyxl.load_workbook(out_path, data_only=True)
    ws2 = wb["Sheet2"]
    wsw = wb["希望"]
    # 医師列 index（1-based）
    hdr2 = [c.value for c in ws2[1]]
    col = {name: i + 1 for i, name in enumerate(hdr2)}
    # date→row（両シート同構造）
    rows = {}
    for r in range(2, ws2.max_row + 1):
        v = ws2.cell(row=r, column=1).value
        if v is not None:
            rows[pd.Timestamp(v).date()] = r
    return wb, ws2, wsw, col, rows


def test_build_end_to_end(tmp_path):
    inp = build_input(tmp_path / "in.xlsx")
    form = write_responses(tmp_path / "resp.csv")
    out = str(tmp_path / "out.xlsx")
    res = bwf.build(form, inp, out)

    wb, ws2, wsw, col, rows = _read(out)
    d2 = datetime.date(2026, 3, 2)
    d9 = datetime.date(2026, 3, 9)
    d16 = datetime.date(2026, 3, 16)
    d23 = datetime.date(2026, 3, 23)
    d3 = datetime.date(2026, 3, 3)

    # 医A: できない 3/2・3/9 → Sheet2=0
    assert ws2.cell(row=rows[d2], column=col["医A"]).value == 0
    assert ws2.cell(row=rows[d9], column=col["医A"]).value == 0
    # 3/2 は不可が優先 → 希望は空
    assert wsw.cell(row=rows[d2], column=col["医A"]).value in (None, "")
    # 3/16 は避け → ×2
    assert wsw.cell(row=rows[d16], column=col["医A"]).value == "×2"
    # 3/23 はやりたい → ○2
    assert wsw.cell(row=rows[d23], column=col["医A"]).value == "○2"
    # 医B: 3/3 避け → ×2
    assert wsw.cell(row=rows[d3], column=col["医B"]).value == "×2"
    wb.close()

    # 集計・警告
    assert res["n_cant"] == 2 and res["n_avoid"] == 2 and res["n_want"] == 1
    assert "医不在" in res["unmatched_names"]
    assert "医C" in res["non_responders"]      # 未回答
    assert set(res["responders"]) == {"医A", "医B"}
    assert res["unknown_labels"] == []


def test_build_reuses_existing_wish_sheet(tmp_path):
    inp = build_input(tmp_path / "in2.xlsx", with_wish=True)
    form = write_responses(tmp_path / "resp2.csv")
    out = str(tmp_path / "out2.xlsx")
    bwf.build(form, inp, out)
    wb = openpyxl.load_workbook(out)
    # 「希望」シートは1枚のまま（重複作成しない）
    assert wb.sheetnames.count("希望") == 1
    wb.close()


def test_build_stage_override(tmp_path):
    inp = build_input(tmp_path / "in3.xlsx")
    form = write_responses(tmp_path / "resp3.csv")
    out = str(tmp_path / "out3.xlsx")
    bwf.build(form, inp, out, avoid_stage=3, want_stage=1)
    wb, ws2, wsw, col, rows = _read(out)
    assert wsw.cell(row=rows[datetime.date(2026, 3, 16)], column=col["医A"]).value == "×3"
    assert wsw.cell(row=rows[datetime.date(2026, 3, 23)], column=col["医A"]).value == "○1"
    wb.close()


def test_build_missing_name_col_raises(tmp_path):
    inp = build_input(tmp_path / "in4.xlsx")
    bad = str(tmp_path / "bad.csv")
    pd.DataFrame([{"foo": "医A", "bar": "3/2"}]).to_csv(bad, index=False)
    with pytest.raises(ValueError):
        bwf.build(bad, inp, str(tmp_path / "o.xlsx"))


def test_roundtrip_converter_output_read_by_solver(tmp_path):
    """変換出力の「希望」シートを solver_cpsat がそのまま読める（Form→入力→ソルバー）。"""
    import solver_cpsat  # noqa: E402
    inp = build_input(tmp_path / "in5.xlsx")
    form = write_responses(tmp_path / "resp5.csv")
    out = str(tmp_path / "out5.xlsx")
    bwf.build(form, inp, out)
    data = solver_cpsat.InputData(out, verbose=False)
    d16 = pd.Timestamp(2026, 3, 16)
    d23 = pd.Timestamp(2026, 3, 23)
    d3 = pd.Timestamp(2026, 3, 3)
    assert data.wish_marks["医A"][d16] == -2   # 避け ×2 → -2
    assert data.wish_marks["医A"][d23] == 2    # やりたい ○2 → +2
    assert data.wish_marks["医B"][d3] == -2
    # 3/2 は不可優先で希望に載っていない
    assert pd.Timestamp(2026, 3, 2) not in data.wish_marks.get("医A", {})
