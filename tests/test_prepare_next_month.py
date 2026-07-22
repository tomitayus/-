# -*- coding: utf-8 -*-
"""prepare_next_month.py の独立テスト。

既存 fixture（conftest の run_output 等）には依存しない。合成xlsxを
tmp_path 上に構築して prepare() を直接検証する。

検証項目:
- 累計加算の検算（旧累計 + 当月実績 = 新累計）
- 氏名照合は normalize + 完全一致のみ（同姓 佐藤彰/佐藤悠/佐藤 の混同なし）
- 名簿に無い医師は加算スキップ + unmatched 警告
- UNASSIGNED / マーカー文字列は医師として集計しない
- 平日でも休日枠扱いの列（C,D,F,G = 0-based 2,3,5,6）の休日分類
- 翌月グリッド生成（日数・祝日反映・カテ当番空欄）
- 祝日サンプルが当月に無い場合のフォールバック（祝日昼/夜 ← 日曜昼/夜）
- Sheet2 の全マーク消去 + 日付の翌月化
- ドラフト警告シート
"""

import datetime
import os
import sys

import openpyxl
import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

import prepare_next_month as pnm

# 実雛形と同じ26列レイアウト（0-based: 0=Date, 1-24=病院, 25=カテ当番）
HEADERS = [
    "Date", "大学平日", "大学土曜昼", "大学土曜夜", "大学日曜昼", "大学日曜夜",
    "大学祝日昼", "大学祝日夜", "支援平日", "支援日直", "支援当直", "夜間急病",
    "太陽の国", "福島南循環器", "大原医療", "星", "しのぶ", "熱海", "日赤",
    "相馬中央", "須賀川", "済生会第一", "済生会第三", "太田西ノ内", "ふたば医療",
    "カテ当番",
]
CUM_COLS = ["全合計", "大学合計", "外病院合計", "平日", "休日合計"]
ROSTER = ["佐藤彰", "佐藤悠", "佐藤勇", "武藤", "小林"]
OLD_CUM = {
    "佐藤彰": [3, 2, 1, 2, 1],
    "佐藤悠": [3, 1, 2, 2, 1],
    "佐藤勇": [4, 1, 3, 3, 1],
    "武藤": [4, 1, 3, 3, 1],
    "小林": [3, 2, 1, 1, 2],
}


def _frame_count(d, j, holidays):
    """合成雛形の枠数ルール（曜日タイプで決定的）。"""
    holi = d in holidays
    wd = d.weekday()
    if j == 1:   # 大学平日
        return 1 if (wd < 5 and not holi) else 0
    if j in (2, 3):  # 大学土曜昼/夜
        return 1 if (wd == 5 and not holi) else 0
    if j in (4, 5):  # 大学日曜昼/夜
        return 1 if (wd == 6 and not holi) else 0
    if j in (6, 7):  # 大学祝日昼/夜
        return 1 if holi else 0
    if j == 8:   # 支援平日
        return 1 if (wd < 5 and not holi) else 0
    if j == 11:  # 夜間急病（毎日）
        return 1
    if j == 24:  # ふたば医療（毎日）
        return 1
    return 0


def _month_dates(year, month):
    import calendar
    n = calendar.monthrange(year, month)[1]
    return [datetime.date(year, month, d) for d in range(1, n + 1)]


def build_template(path, year, month, holidays):
    """実雛形と同構造の合成雛形（sheet1/Sheet2/Sheet3）を作る。"""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "sheet1"
    for c, h in enumerate(HEADERS, start=1):
        ws1.cell(row=1, column=c, value=h)
    for i, d in enumerate(_month_dates(year, month)):
        r = 2 + i
        ws1.cell(row=r, column=1, value=datetime.datetime(d.year, d.month, d.day))
        for j in range(1, 25):
            ws1.cell(row=r, column=j + 1, value=_frame_count(d, j, holidays))
        ws1.cell(row=r, column=26, value="A")  # カテ当番（ドラフトでは消えるべき）

    ws2 = wb.create_sheet("Sheet2")
    ws2.cell(row=1, column=1, value="Date")
    for c, nm in enumerate(ROSTER, start=2):
        ws2.cell(row=1, column=c, value=nm)
    for i, d in enumerate(_month_dates(year, month)):
        r = 2 + i
        ws2.cell(row=r, column=1, value=datetime.datetime(d.year, d.month, d.day))
        ws2.cell(row=r, column=2, value=1)  # 消去されるべきマーク
        ws2.cell(row=r, column=3, value=3)

    ws3 = wb.create_sheet("Sheet3")
    hdr = ["氏名", "属性", "カテ当番", "出張日", "出張先"] + CUM_COLS
    for c, h in enumerate(hdr, start=1):
        ws3.cell(row=1, column=c, value=h)
    for i, nm in enumerate(ROSTER):
        r = 2 + i
        ws3.cell(row=r, column=1, value=nm)
        ws3.cell(row=r, column=2, value=2)
        for k, v in enumerate(OLD_CUM[nm]):
            ws3.cell(row=r, column=6 + k, value=v)
    wb.save(path)


def build_output(path, year, month, assignments):
    """pattern_01 シートを持つ合成出力を作る（1行目バナー・2行目ヘッダ）。

    assignments: {(day, col_j): 値} 。値は医師名 or 任意文字列。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "pattern_01"
    ws.cell(row=1, column=1, value="✅ 絶対禁忌チェック全クリア（ダミーバナー）")
    for c, h in enumerate(HEADERS, start=1):
        ws.cell(row=2, column=c, value=h)
    wd_jp = ["月", "火", "水", "木", "金", "土", "日"]
    for i, d in enumerate(_month_dates(year, month)):
        r = 3 + i
        ws.cell(row=r, column=1,
                value=f"{d.year}/{d.month}/{d.day} ({wd_jp[d.weekday()]})")
        for (day, j), v in assignments.items():
            if day == d.day:
                ws.cell(row=r, column=j + 1, value=v)
    wb.save(path)


# ---------------------------------------------------------------------------
# メインシナリオ: 2026年3月 → 4月ドラフト
# ---------------------------------------------------------------------------
MAR_HOLIDAYS = {datetime.date(2026, 3, 20)}  # 春分の日（金）

# 3月の割当（day, 0-based列j）→ 値
# 2026-03: 1日=日, 2日=月, 4日=水, 7日=土, 20日=金祝
ASSIGN = {
    (2, 1): "佐藤彰",    # 月・大学平日        → bg, 平日
    (7, 11): "佐藤彰",   # 土・夜間急病        → ht, 休日
    (4, 2): "佐藤悠",    # 水・大学土曜昼(列2) → bg, 平日でも休日扱い
    (20, 6): "佐藤勇",   # 金祝・大学祝日昼    → bg, 休日
    (3, 11): "佐藤",     # 名簿に無い（部分一致してはならない）
    (5, 11): "UNASSIGNED",  # 医師ではない
    (6, 11): "〇",          # マーカーは医師ではない
    (10, 24): "小林",    # 火・ふたば医療      → ht, 平日
}
# 期待実績 [全, 大学, 外, 平日, 休日]
EXPECT_ADD = {
    "佐藤彰": [2, 1, 1, 1, 1],
    "佐藤悠": [1, 1, 0, 0, 1],
    "佐藤勇": [1, 1, 0, 0, 1],
    "武藤": [0, 0, 0, 0, 0],
    "小林": [1, 0, 1, 1, 0],
}


@pytest.fixture(scope="module")
def draft(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("pnm")
    tpl = str(tmp / "template_2603.xlsx")
    out = str(tmp / "output_2603.xlsx")
    dst = str(tmp / "draft_2604.xlsx")
    build_template(tpl, 2026, 3, MAR_HOLIDAYS)
    build_output(out, 2026, 3, ASSIGN)
    res = pnm.prepare(out, tpl, 1, dst)
    wb = openpyxl.load_workbook(dst)
    return res, wb


def test_cumulative_addition_exact(draft):
    """旧累計 + 当月実績 = 新累計（全医師×5列の機械照合）。"""
    res, wb = draft
    ws3 = wb["Sheet3"]
    hdr = {ws3.cell(row=1, column=c).value: c
           for c in range(1, ws3.max_column + 1)}
    got = {}
    for r in range(2, ws3.max_row + 1):
        nm = ws3.cell(row=r, column=1).value
        if nm:
            got[nm] = [ws3.cell(row=r, column=hdr[k]).value for k in CUM_COLS]
    for nm in ROSTER:
        expect = [o + a for o, a in zip(OLD_CUM[nm], EXPECT_ADD[nm])]
        assert got[nm] == expect, f"{nm}: 期待{expect} 実際{got[nm]}"


def test_exact_match_only_no_fuzzy(draft):
    """『佐藤』は佐藤彰/佐藤悠/佐藤勇のどれにも吸収されず unmatched になる。"""
    res, _ = draft
    assert "佐藤" in res["unmatched"]
    # 同姓3人はそれぞれ正しく照合済み（unmatchedに入らない）
    for nm in ("佐藤彰", "佐藤悠", "佐藤勇"):
        assert nm not in res["unmatched"]


def test_non_doctor_strings_ignored(draft):
    """UNASSIGNED・〇 は医師として集計されない。"""
    res, _ = draft
    assert "UNASSIGNED" not in res["month_stats"]
    assert "〇" not in res["month_stats"]
    assert "UNASSIGNED" not in res["unmatched"]


def test_next_month_grid(draft):
    """4月グリッド: 30日連続・4/29祝日反映・カテ当番空欄。"""
    _, wb = draft
    ws1 = wb["sheet1"]
    dates = []
    r = 2
    while True:
        v = ws1.cell(row=r, column=1).value
        if not isinstance(v, (datetime.datetime, datetime.date)):
            break
        dates.append(v.date() if isinstance(v, datetime.datetime) else v)
        r += 1
    assert dates == _month_dates(2026, 4)
    # 4/29（水・昭和の日）= 行30
    r429 = 2 + 28
    assert ws1.cell(row=r429, column=2).value == 0      # 大学平日
    assert ws1.cell(row=r429, column=7).value == 1      # 大学祝日昼
    assert ws1.cell(row=r429, column=8).value == 1      # 大学祝日夜
    # 4/1（水・平日）
    assert ws1.cell(row=2, column=2).value == 1         # 大学平日
    assert ws1.cell(row=2, column=7).value == 0         # 大学祝日昼
    # カテ当番列は全行空欄
    for rr in range(2, 2 + 30):
        assert ws1.cell(row=rr, column=26).value in (None, "")
    # 3月31日の残骸行が無い
    assert ws1.cell(row=2 + 30, column=1).value is None


def test_weekday_saturday_sunday_patterns(draft):
    """曜日タイプ別の枠数が当月ルール通りに再現される。"""
    _, wb = draft
    ws1 = wb["sheet1"]
    for i, d in enumerate(_month_dates(2026, 4)):
        r = 2 + i
        holi = (d == datetime.date(2026, 4, 29))
        for j in (1, 2, 4, 8, 11, 24):
            expect = _frame_count(d, j, {d} if holi else set())
            got = ws1.cell(row=r, column=j + 1).value
            assert got == expect, f"{d} col{j}: 期待{expect} 実際{got}"


def test_sheet2_blank_and_next_month(draft):
    """Sheet2: マーク全消去・日付は4月。"""
    _, wb = draft
    ws2 = wb["Sheet2"]
    dates = []
    marks = 0
    for r in range(2, ws2.max_row + 1):
        v = ws2.cell(row=r, column=1).value
        if isinstance(v, (datetime.datetime, datetime.date)):
            dates.append(v.date() if isinstance(v, datetime.datetime) else v)
        for c in range(2, ws2.max_column + 1):
            if ws2.cell(row=r, column=c).value is not None:
                marks += 1
    assert dates == _month_dates(2026, 4)
    assert marks == 0


def test_draft_warning_sheet(draft):
    """先頭シートにドラフト警告が明示される。"""
    _, wb = draft
    first = wb.worksheets[0]
    assert "ドラフト" in str(first["A1"].value)
    assert "手動確認" in str(first["A1"].value)


def test_result_metadata(draft):
    res, _ = draft
    assert res["cur_ym"] == (2026, 3)
    assert res["nxt_ym"] == (2026, 4)
    assert datetime.date(2026, 4, 29) in res["nxt_holidays"]


# ---------------------------------------------------------------------------
# 祝日フォールバック: 当月に祝日が無い月（2026年6月 → 7月、7/20海の日）
# ---------------------------------------------------------------------------
def test_holiday_fallback_without_sample(tmp_path):
    """当月(6月)に祝日サンプルが無くても、翌月の祝日枠は日曜推定を流用する。"""
    tpl = str(tmp_path / "template_2606.xlsx")
    out = str(tmp_path / "output_2606.xlsx")
    dst = str(tmp_path / "draft_2607.xlsx")
    build_template(tpl, 2026, 6, set())  # 6月は祝日なし
    build_output(out, 2026, 6, {(1, 1): "小林"})  # 6/1(月) 大学平日
    res = pnm.prepare(out, tpl, 1, dst)
    assert res["nxt_ym"] == (2026, 7)
    assert datetime.date(2026, 7, 20) in res["nxt_holidays"]  # 海の日
    wb = openpyxl.load_workbook(dst)
    ws1 = wb["sheet1"]
    r720 = 2 + 19  # 7/20
    v = ws1.cell(row=r720, column=1).value
    assert v.date() == datetime.date(2026, 7, 20)
    # 祝日昼/夜 ← 日曜昼/夜のフォールバック（合成ルールでは日曜=1）
    assert ws1.cell(row=r720, column=7).value == 1       # 大学祝日昼 = 日曜昼の推定値
    assert ws1.cell(row=r720, column=8).value == 1       # 大学祝日夜 = 日曜夜の推定値
    assert ws1.cell(row=r720, column=2).value == 0       # 大学平日 = 0
    assert ws1.cell(row=r720, column=5).value == 0       # 大学日曜昼も0（祝日行）


# ---------------------------------------------------------------------------
# 異常系
# ---------------------------------------------------------------------------
def test_missing_pattern_sheet_raises(tmp_path):
    tpl = str(tmp_path / "t.xlsx")
    out = str(tmp_path / "o.xlsx")
    build_template(tpl, 2026, 3, MAR_HOLIDAYS)
    build_output(out, 2026, 3, {})
    with pytest.raises(ValueError, match="pattern_05"):
        pnm.prepare(out, tpl, 5, str(tmp_path / "d.xlsx"))
