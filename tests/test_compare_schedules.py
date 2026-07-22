# -*- coding: utf-8 -*-
"""compare_schedules.py の独立テスト。

合成 pattern ワークブックを tmp_path 上に構築し、突合ロジックを検証する。
グリッド解釈は prepare_next_month に一本化しているため、分類（BG/HT）や
氏名正規化の挙動は prepare_next_month のテストと整合する前提。

検証項目:
- 同一グリッド → セル相違0・回数差0・一致率100%
- 1セル差し替え → 相違1件・該当2医師の Δ合計が ±1
- 列位置による分類（大学 B-K / 外 L-Y）が回数差に正しく反映
- UNASSIGNED / 数値 / 空欄セルは医師として集計しない
- 片方にしか無い日付は日付カバレッジ不一致として拾う
- pattern シート名の桁揺れ（pattern_1 / pattern_01）を解決
- カテ当番(Z)列は突合対象から除外
"""

import datetime
import os
import sys

import openpyxl
import pytest

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if REPO_ROOT not in sys.path:
    sys.path.insert(0, REPO_ROOT)

import compare_schedules as cs

# 実雛形と同じ26列レイアウト（0-based: 0=Date, 1-24=病院, 25=カテ当番）
HEADERS = [
    "Date", "大学平日", "大学土曜昼", "大学土曜夜", "大学日曜昼", "大学日曜夜",
    "大学祝日昼", "大学祝日夜", "支援平日", "支援日直", "支援当直", "夜間急病",
    "太陽の国", "福島南循環器", "大原医療", "星", "しのぶ", "熱海", "日赤",
    "相馬中央", "須賀川", "済生会第一", "済生会第三", "太田西ノ内", "ふたば医療",
    "カテ当番",
]


def build_pattern(path, assignments, sheet_name="pattern_01", year=2026, month=8):
    """合成 pattern ワークブックを作る。

    assignments: {day(int) -> {j(0-based grid index) -> セル値}}
      j=1 が大学平日(BG), j=11 が夜間急病(HT)。値は医師名 or 数値/マーカー。
    行1にヘッダ（A1="Date"）、行2以降に日付＋割当を書く。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for c, h in enumerate(HEADERS, start=1):
        ws.cell(row=1, column=c, value=h)
    days = sorted(assignments.keys())
    for i, day in enumerate(days):
        r = 2 + i
        d = datetime.date(year, month, day)
        ws.cell(row=r, column=1, value=datetime.datetime(d.year, d.month, d.day))
        for j, val in assignments[day].items():
            ws.cell(row=r, column=j + 1, value=val)
    wb.save(path)
    return path


# ---------------------------------------------------------------------------
def test_identical_grids_no_diff(tmp_path):
    a = build_pattern(tmp_path / "a.xlsx",
                      {1: {1: "小林", 11: "武藤"}, 2: {1: "佐藤彰", 11: "小林"}})
    b = build_pattern(tmp_path / "b.xlsx",
                      {1: {1: "小林", 11: "武藤"}, 2: {1: "佐藤彰", 11: "小林"}})
    res = cs.compare(str(a), str(b))
    assert res["cell_diffs"] == []
    assert all(r["d_total"] == 0 and r["d_bg"] == 0 and r["d_ht"] == 0
               for r in res["count_rows"])
    assert "一致率: **100.0%**" in res["report"]


def test_single_swap(tmp_path):
    # 2日目の大学平日(j=1)を 佐藤彰 → 武藤 に差し替え
    a = build_pattern(tmp_path / "a.xlsx",
                      {1: {1: "小林", 11: "武藤"}, 2: {1: "佐藤彰", 11: "小林"}})
    b = build_pattern(tmp_path / "b.xlsx",
                      {1: {1: "小林", 11: "武藤"}, 2: {1: "武藤", 11: "小林"}})
    res = cs.compare(str(a), str(b))
    assert len(res["cell_diffs"]) == 1
    diff = res["cell_diffs"][0]
    assert diff["date"] == datetime.date(2026, 8, 2)
    assert diff["col_label"] == "大学平日"
    assert diff["a"] == "佐藤彰" and diff["b"] == "武藤"

    by = {r["norm"]: r for r in res["count_rows"]}
    # A では佐藤彰が大学枠に居て B では居ない → Δ合計 +1, Δ大学 +1
    assert by["佐藤彰"]["d_total"] == 1 and by["佐藤彰"]["d_bg"] == 1
    # 武藤は A で1回・B で2回 → Δ合計 -1, Δ大学 -1
    assert by["武藤"]["d_total"] == -1 and by["武藤"]["d_bg"] == -1


def test_bg_vs_ht_classification(tmp_path):
    # 小林: A=大学(j=1)、B=外(j=11)。合計は同じでも大学/外の内訳が動く
    a = build_pattern(tmp_path / "a.xlsx", {1: {1: "小林"}})
    b = build_pattern(tmp_path / "b.xlsx", {1: {11: "小林"}})
    res = cs.compare(str(a), str(b))
    by = {r["norm"]: r for r in res["count_rows"]}
    assert by["小林"]["d_total"] == 0
    assert by["小林"]["d_bg"] == 1   # A に大学1
    assert by["小林"]["d_ht"] == -1  # B に外1
    # 同じ枠(j=1) 上では A=小林 / B=空 で相違、j=11 も相違 → 2件
    assert len(res["cell_diffs"]) == 2


def test_markers_and_numbers_skipped(tmp_path):
    a = build_pattern(tmp_path / "a.xlsx",
                      {1: {1: "小林", 2: "UNASSIGNED", 11: 1}})
    b = build_pattern(tmp_path / "b.xlsx", {1: {1: "小林"}})
    res = cs.compare(str(a), str(b))
    # UNASSIGNED も数値1も医師ではない → 相違なし・回数差なし
    assert res["cell_diffs"] == []
    by = {r["norm"]: r for r in res["count_rows"]}
    assert by["小林"]["d_total"] == 0


def test_date_coverage_mismatch(tmp_path):
    a = build_pattern(tmp_path / "a.xlsx", {1: {1: "小林"}, 2: {1: "武藤"}})
    b = build_pattern(tmp_path / "b.xlsx", {1: {1: "小林"}})
    res = cs.compare(str(a), str(b))
    assert "日付カバレッジ不一致" in res["report"]
    # 2日目は A のみ割当 → 相違1件（B側は空）
    assert len(res["cell_diffs"]) == 1
    assert res["cell_diffs"][0]["a"] == "武藤" and res["cell_diffs"][0]["b"] == "—"


def test_pattern_name_digit_variant(tmp_path):
    a = build_pattern(tmp_path / "a.xlsx", {1: {1: "小林"}}, sheet_name="pattern_1")
    b = build_pattern(tmp_path / "b.xlsx", {1: {1: "小林"}}, sheet_name="pattern_01")
    res = cs.compare(str(a), str(b))  # --pattern 1 で両方解決できる
    assert res["cell_diffs"] == []


def test_kate_column_excluded(tmp_path):
    # カテ当番(j=25)に文字列が入っていても突合対象外
    a = build_pattern(tmp_path / "a.xlsx", {1: {1: "小林", 25: "A"}})
    b = build_pattern(tmp_path / "b.xlsx", {1: {1: "小林", 25: "B"}})
    res = cs.compare(str(a), str(b))
    assert res["cell_diffs"] == []
    by = {r["norm"]: r for r in res["count_rows"]}
    assert "A" not in by and "B" not in by


def test_missing_pattern_sheet_raises(tmp_path):
    p = build_pattern(tmp_path / "a.xlsx", {1: {1: "小林"}}, sheet_name="pattern_01")
    with pytest.raises(ValueError):
        cs.load_grid(str(p), pattern_num=5)
